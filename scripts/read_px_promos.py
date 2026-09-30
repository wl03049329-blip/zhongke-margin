"""Read-only workbook extraction; dynamic header/period discovery."""
import sys, json, re, unicodedata
from datetime import date
from openpyxl import load_workbook

def clean(value):
    return re.sub(r'\s+', ' ', unicodedata.normalize('NFKC', str(value or ''))).strip()

workbook = load_workbook(sys.argv[1], data_only=True)
output = {'rows': [], 'issues': []}
for sheet in workbook:
    header = next((r for r in sheet.iter_rows() if any('商品' in clean(c.value) and '代號' in clean(c.value) for c in r)), None)
    if not header:
        output['issues'].append({'severity':'error','sheet':sheet.title,'reason':'HEADER_NOT_FOUND'})
        continue
    fields = {}
    for cell in header:
        text = clean(cell.value).replace(' ', '')
        for key, needle in [('code','商品代號'),('barcode','條碼'),('name','品名')]:
            if needle in text: fields[key] = cell.column - 1
    top = ' '.join(clean(c.value) for row in sheet.iter_rows(max_row=header[0].row) for c in row)
    year_match = re.search(r'(20\d{2})', top + ' ' + sheet.title)
    year = int(sys.argv[2]) if len(sys.argv)>2 else int(year_match[1]) if year_match else None
    periods = []
    for cell in header:
        text = clean(cell.value)
        m = re.search(r'(\d{1,2})-(\d+)檔.*?(\d{1,2})/(\d{1,2})\s*[-~～]\s*(\d{1,2})/(\d{1,2})',text)
        if m and year:
            try:
                start = date(year,int(m[3]),int(m[4])); end = date(year+(int(m[5])<int(m[3])),int(m[5]),int(m[6]))
                if end < start: raise ValueError('date order')
                periods.append({'column':cell.column-1,'id':f'{year}-{int(m[1]):02d}-{m[2]}','label':f'{int(m[1])}-{m[2]}','startDate':start.isoformat(),'endDate':end.isoformat()})
            except ValueError: output['issues'].append({'severity':'error','sheet':sheet.title,'reason':'INVALID_DATE','raw':text})
    if len(fields)!=3 or not periods:
        output['issues'].append({'severity':'error','sheet':sheet.title,'reason':'MISSING_FIELDS_OR_PERIODS'})
        continue
    for row in sheet.iter_rows(min_row=header[0].row+1):
        raw = {key: row[col].value for key,col in fields.items()}
        if not any(raw.values()):
            output['issues'].append({'severity':'warning','sheet':sheet.title,'row':row[0].row,'reason':'BLANK_ROW'})
            continue
        output['rows'].append({'sheet':sheet.title,'row':row[0].row,'raw':raw,'periods':[{**p,'rawText':row[p['column']].value} for p in periods]})
print(json.dumps(output,ensure_ascii=False,default=str))
